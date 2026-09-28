"""Converte o HTML de slides em uma apresentação PPTX.

O HTML é renderizado pelo mesmo Chromium usado na exportação para PDF. Cada
página do PDF resultante vira um slide, preservando com fidelidade o layout,
as fontes, as imagens e os estilos do material exibido no editor.
"""

from io import BytesIO
from pathlib import Path
import tempfile

import pymupdf
from pptx import Presentation
from pptx.util import Inches

from graph.tools.sandbox.shared.html_pdf_tools import render_html_to_pdf


SLIDE_WIDTH_INCHES = 11.69
SLIDE_HEIGHT_INCHES = 8.27
RASTER_SCALE = 2.0


class HtmlToPptxError(RuntimeError):
    """Erro ao montar a apresentação a partir do HTML renderizado."""


def convert_html_bytes_to_pptx(
    content: bytes,
    title: str = "apresentacao",
    work_dir: str | Path | None = None,
) -> bytes:
    """Renderiza um HTML horizontal e retorna um PPTX com um slide por página."""
    if not content.strip():
        raise HtmlToPptxError("O HTML não pode estar vazio")

    temporary_parent = Path(work_dir).resolve() if work_dir else None
    if temporary_parent:
        temporary_parent.mkdir(parents=True, exist_ok=True)

    with tempfile.TemporaryDirectory(prefix="html-to-pptx-", dir=temporary_parent) as directory:
        temp_dir = Path(directory)
        pdf_path = temp_dir / "slides.pdf"
        render_html_to_pdf(content, pdf_path, "H", temp_dir)

        try:
            document = pymupdf.open(pdf_path)
        except (pymupdf.FileDataError, OSError) as exc:
            raise HtmlToPptxError("Não foi possível abrir os slides renderizados") from exc

        try:
            if document.page_count == 0:
                raise HtmlToPptxError("O HTML não gerou nenhum slide")

            presentation = Presentation()
            presentation.slide_width = Inches(SLIDE_WIDTH_INCHES)
            presentation.slide_height = Inches(SLIDE_HEIGHT_INCHES)
            presentation.core_properties.title = title

            blank_layout = presentation.slide_layouts[6]
            matrix = pymupdf.Matrix(RASTER_SCALE, RASTER_SCALE)
            for page in document:
                pixmap = page.get_pixmap(matrix=matrix, alpha=False)
                image = BytesIO(pixmap.tobytes("png"))
                slide = presentation.slides.add_slide(blank_layout)
                slide.shapes.add_picture(
                    image,
                    0,
                    0,
                    width=presentation.slide_width,
                    height=presentation.slide_height,
                )

            output = BytesIO()
            presentation.save(output)
            return output.getvalue()
        except HtmlToPptxError:
            raise
        except (ValueError, OSError) as exc:
            raise HtmlToPptxError("Não foi possível gerar a apresentação PPTX") from exc
        finally:
            document.close()
