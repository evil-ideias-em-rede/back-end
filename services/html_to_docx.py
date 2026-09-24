"""Converte o HTML A4 gerado por ``pdf_to_html.py`` para DOCX.

O conversor mantém cada folha A4 como uma página do Word. A imagem da folha
fica ancorada atrás do texto e os textos extraídos do HTML continuam editáveis.
"""

from __future__ import annotations

import base64
import io
import re
import tempfile
from pathlib import Path

from bs4 import BeautifulSoup, NavigableString, Tag
from docx import Document
from docx.enum.text import WD_BREAK
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls, qn
from docx.shared import Emu, Mm, Pt, RGBColor
from PIL import Image


A4_WIDTH_MM = 210.0
A4_HEIGHT_MM = 297.0
TWIPS_PER_PX = 1440 / 96
EMU_PER_PX = 914400 / 96
PT_PER_PX = 72 / 96
NORMAL_LINE_HEIGHT = 1.15
NUMBER = r"(-?\d+(?:\.\d+)?)"

FONT_MAP = {
    "helveticaneue": "Arial", "helvetica": "Arial", "arial": "Arial",
    "arialmt": "Arial", "timesnewroman": "Times New Roman",
    "timesnewromanps": "Times New Roman", "timesnewromanpsmt": "Times New Roman",
    "times": "Times New Roman", "timesroman": "Times New Roman",
    "courier": "Courier New", "couriernew": "Courier New", "calibri": "Calibri",
    "cambria": "Cambria", "georgia": "Georgia", "verdana": "Verdana",
    "tahoma": "Tahoma", "garamond": "Garamond", "symbol": "Symbol",
}


class HtmlToDocxError(RuntimeError):
    """Erro de conversão HTML para DOCX."""


def _clean_font(raw: str) -> tuple[str, bool, bool]:
    name = raw.strip().strip("'\"")
    name = re.sub(r"^[A-Z]{6}\+", "", name)
    base, _, style = name.partition("-")
    style = style.lower()
    bold = any(value in style for value in ("bold", "black", "heavy", "semibold"))
    italic = any(value in style for value in ("italic", "oblique"))
    key = base.lower().replace(" ", "")
    if key in FONT_MAP:
        return FONT_MAP[key], bold, italic
    for suffix in ("psmt", "mt", "ps"):
        if key.endswith(suffix) and key[:-len(suffix)] in FONT_MAP:
            return FONT_MAP[key[:-len(suffix)]], bold, italic
    readable = re.sub(r"(?<=[a-z])(?=[A-Z])", " ", re.sub(r"(PSMT|MT|PS)$", "", base))
    return readable or "Arial", bold, italic


def _read_css(html_text: str):
    fonts = {}
    for name, body in re.findall(r"\.(ft\d+)\s*\{([^}]*)\}", html_text):
        props = {
            key.strip().lower(): value.strip()
            for key, _, value in (part.partition(":") for part in body.split(";"))
            if key.strip()
        }
        size = re.match(NUMBER, props.get("font-size", "16px"))
        line_height = re.match(NUMBER, props.get("line-height", ""))
        family, bold, italic = _clean_font(props.get("font-family", "Arial"))
        color = props.get("color", "#000000").lstrip("#")
        if len(color) == 3:
            color = "".join(char * 2 for char in color)
        fonts[name] = {
            "size": float(size.group(1)) if size else 16.0,
            "line_height": float(line_height.group(1)) if line_height else None,
            "family": family,
            "bold": bold or "bold" in props.get("font-weight", ""),
            "italic": italic or "italic" in props.get("font-style", ""),
            "color": color if re.fullmatch(r"[0-9a-fA-F]{6}", color) else "000000",
        }

    pages = {}
    for number, width, height, scale in re.findall(
        rf"\.pg(\d+)-conteudo\{{width:{NUMBER}px;height:{NUMBER}px;transform:scale\({NUMBER}\);\}}",
        html_text,
    ):
        pages[number] = {"width": float(width), "slice_height": float(height), "scale": float(scale)}
    for number, url, width, height in re.findall(
        rf"\.pg(\d+)-fundo\{{background-image:url\(\"(data:[^\"]+)\"\);background-size:{NUMBER}px {NUMBER}px;\}}",
        html_text,
    ):
        pages.setdefault(number, {}).update({"background": url, "background_width": float(width), "background_height": float(height)})
    return fonts, pages


def _decode_data_url(url: str) -> Image.Image:
    _, _, encoded = url.partition(",")
    return Image.open(io.BytesIO(base64.b64decode(encoded))).convert("RGB")


def _configure_a4(document: Document) -> None:
    section = document.sections[0]
    section.page_width = Mm(A4_WIDTH_MM)
    section.page_height = Mm(A4_HEIGHT_MM)
    for attr in ("left_margin", "right_margin", "top_margin", "bottom_margin", "header_distance", "footer_distance", "gutter"):
        setattr(section, attr, Emu(0))
    style = document.styles["Normal"]
    style.font.name = "Arial"
    style.font.size = Pt(1)
    style.paragraph_format.space_before = style.paragraph_format.space_after = Pt(0)
    style.paragraph_format.line_spacing = 1.0


def _anchor_paragraph(document: Document, page_break: bool):
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.space_before = paragraph.paragraph_format.space_after = Pt(0)
    if page_break:
        paragraph.paragraph_format.page_break_before = True
    return paragraph


def _insert_background(paragraph, image: Image.Image, width_emu: int, height_emu: int, index: int) -> None:
    buffer = io.BytesIO()
    image.save(buffer, format="PNG", optimize=True)
    buffer.seek(0)
    run = paragraph.add_run()
    run.add_picture(buffer, width=Emu(width_emu), height=Emu(height_emu))
    inline = run._r.find(qn("w:drawing")).find(qn("wp:inline"))
    graphic = inline.find(qn("a:graphic"))
    anchor = parse_xml(
        f'''<wp:anchor {nsdecls("wp", "a", "pic", "r")} distT="0" distB="0" distL="0" distR="0"
        simplePos="0" relativeHeight="{index}" behindDoc="1" locked="1" layoutInCell="1" allowOverlap="1">
        <wp:simplePos x="0" y="0"/>
        <wp:positionH relativeFrom="page"><wp:posOffset>0</wp:posOffset></wp:positionH>
        <wp:positionV relativeFrom="page"><wp:posOffset>0</wp:posOffset></wp:positionV>
        <wp:extent cx="{width_emu}" cy="{height_emu}"/>
        <wp:effectExtent l="0" t="0" r="0" b="0"/><wp:wrapNone/>
        <wp:docPr id="{1000 + index}" name="Fundo folha {index}"/><wp:cNvGraphicFramePr/>
        </wp:anchor>'''
    )
    anchor.append(graphic)
    inline.getparent().replace(inline, anchor)


def _position_paragraph(paragraph, x_twips: int, y_twips: int) -> None:
    page_width = int(round(A4_WIDTH_MM / 25.4 * 1440))
    width = max(page_width - x_twips, 200)
    ppr = paragraph._p.get_or_add_pPr()
    frame = parse_xml(
        f'<w:framePr {nsdecls("w")} w:w="{width}" w:wrap="around" w:hAnchor="page" '
        f'w:vAnchor="page" w:x="{x_twips}" w:y="{y_twips}" w:hSpace="0" w:vSpace="0"/>'
    )
    ppr.insert(0, frame)
    ppr.append(parse_xml(f'<w:ind {nsdecls("w")} w:left="0" w:right="0" w:firstLine="0"/>'))


def _add_runs(paragraph, node: Tag, font: dict, scale: float, bold=False, italic=False, underline=False):
    for child in node.children:
        if isinstance(child, NavigableString):
            text = str(child).replace("\u00a0", " ")
            if not text:
                continue
            run = paragraph.add_run(text)
            run.font.name = font["family"]
            run._r.get_or_add_rPr().get_or_add_rFonts().set(qn("w:cs"), font["family"])
            run.font.size = Pt(round(font["size"] * scale * PT_PER_PX * 2) / 2)
            run.font.bold = bold or font["bold"]
            run.font.italic = italic or font["italic"]
            run.font.underline = underline or None
            run.font.color.rgb = RGBColor.from_string(font["color"].upper())
        elif isinstance(child, Tag):
            name = child.name.lower()
            if name == "br":
                paragraph.add_run().add_break(WD_BREAK.LINE)
                continue
            _add_runs(
                paragraph,
                child,
                font,
                scale,
                bold or name in ("b", "strong"),
                italic or name in ("i", "em"),
                underline or name == "u",
            )


def convert_html_file_to_docx(html_path: Path, docx_path: Path) -> None:
    html_text = html_path.read_text(encoding="utf-8")
    fonts, pages = _read_css(html_text)
    soup = BeautifulSoup(html_text, "html.parser")
    sheets = soup.select("section.folha")
    if not sheets:
        raise HtmlToDocxError("O HTML não possui folhas A4 compatíveis com o conversor PDF.")

    document = Document()
    _configure_a4(document)
    for paragraph in list(document.paragraphs):
        paragraph._p.getparent().remove(paragraph._p)

    cache: dict[str, Image.Image] = {}
    a4_width_emu = int(Mm(A4_WIDTH_MM))
    total_text = 0
    for index, sheet in enumerate(sheets, start=1):
        number = sheet.get("data-pagina-pdf", "")
        info = pages.get(number)
        content = sheet.select_one(".folha-conteudo")
        if not info or "scale" not in info:
            raise HtmlToDocxError(f"Dados da página {number} não encontrados no CSS.")
        scale = info["scale"]
        anchor = _anchor_paragraph(document, page_break=index > 1)

        if info.get("background") and content is not None:
            match = re.search(rf"background-position:\s*0\s+{NUMBER}px", content.get("style", ""))
            offset = -float(match.group(1)) if match else 0.0
            if number not in cache:
                cache[number] = _decode_data_url(info["background"])
            image = cache[number]
            factor = image.height / info["background_height"]
            top = int(round(offset * factor))
            bottom = int(round(min(offset + info["slice_height"], info["background_height"]) * factor))
            if bottom > top:
                cropped = image.crop((0, top, image.width, bottom))
                css_height = (bottom - top) / factor
                _insert_background(anchor, cropped, a4_width_emu, int(round(css_height * scale * EMU_PER_PX)), index)

        if content is None:
            continue
        for paragraph_html in content.find_all("p", recursive=False):
            if not paragraph_html.get_text().replace("\u00a0", "").strip():
                continue
            style = paragraph_html.get("style", "")
            top_match = re.search(rf"top:{NUMBER}px", style)
            left_match = re.search(rf"left:{NUMBER}px", style)
            top = float(top_match.group(1)) if top_match else 0.0
            left = float(left_match.group(1)) if left_match else 0.0
            class_name = next((value for value in paragraph_html.get("class", []) if value in fonts), None)
            font = fonts.get(class_name, {"size": 16.0, "line_height": None, "family": "Arial", "bold": False, "italic": False, "color": "000000"})
            paragraph = document.add_paragraph()
            _position_paragraph(paragraph, int(round(left * scale * TWIPS_PER_PX)), int(round(top * scale * TWIPS_PER_PX)))
            line_height = font["line_height"] or font["size"] * NORMAL_LINE_HEIGHT
            paragraph.paragraph_format.space_before = paragraph.paragraph_format.space_after = Pt(0)
            paragraph.paragraph_format.line_spacing = Pt(line_height * scale * PT_PER_PX)
            paragraph.paragraph_format.line_spacing_rule = 4
            _add_runs(paragraph, paragraph_html, font, scale)
            total_text += 1
        _anchor_paragraph(document, page_break=False)

    document.core_properties.title = soup.title.get_text(strip=True) if soup.title else html_path.stem
    document.save(str(docx_path))


def convert_html_bytes_to_docx(content: bytes, title: str = "material") -> bytes:
    """Converte bytes HTML e devolve o DOCX pronto para a resposta HTTP."""
    with tempfile.TemporaryDirectory(prefix="html-to-docx-") as directory:
        html_path = Path(directory) / f"{Path(title).stem or 'material'}.html"
        docx_path = Path(directory) / "material.docx"
        html_path.write_bytes(content)
        try:
            convert_html_file_to_docx(html_path, docx_path)
        except HtmlToDocxError:
            raise
        except Exception as exc:
            raise HtmlToDocxError(f"Falha ao gerar o DOCX: {exc}") from exc
        return docx_path.read_bytes()


if __name__ == "__main__":
    import argparse

    parser = argparse.ArgumentParser(description="Converte HTML A4 gerado pelo pdf_to_html.py em DOCX editável.")
    parser.add_argument("html", type=Path, help="Arquivo HTML de entrada")
    parser.add_argument("docx", type=Path, help="Arquivo DOCX de saída")
    args = parser.parse_args()
    convert_html_file_to_docx(args.html, args.docx)
